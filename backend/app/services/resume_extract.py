"""Deterministic resume text extraction + structuring (import-safe).

This module has NO FastAPI/motor imports so it can run inside the isolated
parser child process (`_parse_worker.py`). Both parent and child import from
here; only `resume_parse.py` spawns the subprocess.
"""

from __future__ import annotations

import io
import re
import zipfile

# Skill lexicon (matched case-insensitively, reported in canonical form).
SKILLS = [
    "Java", "Python", "React", "Spring Boot", "SQL", "AWS",
    "JavaScript", "TypeScript", "Node.js", "Express", "NestJS", "Next.js",
    "Angular", "Vue", "HTML", "CSS", "Django", "Flask", "FastAPI",
    "Go", "Rust", "C++", "C#", ".NET", "Kotlin", "Swift", "PHP", "Ruby on Rails",
    "PostgreSQL", "MySQL", "MongoDB", "Redis", "Elasticsearch",
    "Docker", "Kubernetes", "Terraform", "Jenkins", "Git", "Linux",
    "Kafka", "RabbitMQ", "GraphQL", "REST", "Microservices",
    "Machine Learning", "TensorFlow", "PyTorch", "Pandas", "NumPy",
    "Selenium", "Cypress", "JUnit", "Mockito", "Hibernate", "JPA", "Maven", "Gradle",
    "Azure", "GCP", "Snowflake", "Airflow", "Spark", "Figma",
]

LANGUAGES = [
    "Java", "Python", "JavaScript", "TypeScript", "Go", "Rust", "C++",
    "C#", "Kotlin", "Swift", "PHP", "Ruby", "Scala", "R",
]

# Canonical skill -> extra spellings seen in the wild (resumes, job posts).
# Matching is case-insensitive with word boundaries on every variant, so
# "Java" never matches "JavaScript" and "SQL" never matches "NoSQL".
ALIASES: dict[str, list[str]] = {
    "JavaScript": ["js", "javascript"],
    "TypeScript": ["ts", "typescript"],
    "Node.js": ["nodejs", "node.js", "node"],
    "Next.js": ["nextjs", "next.js"],
    "PostgreSQL": ["postgresql", "postgres"],
    "Kubernetes": ["kubernetes", "k8s"],
    ".NET": [".net", "dotnet"],
    "C#": ["c#", "csharp"],
    "C++": ["c++"],
    "React": ["react"],
    "React Native": ["react native", "react-native"],
    "SQL": ["sql"],
    "Go": ["go", "golang"],
    "R": ["r"],
    "C": ["c"],
    "Ruby on Rails": ["ruby on rails", "rails"],
    "Machine Learning": ["machine learning", "ml"],
}


def skill_hit(text: str, skill: str) -> bool:
    """True if `skill` (or a known alias) appears as a standalone token.

    Uses lookarounds instead of `\\b` so symbols like C++/C#/.NET match.
    """
    variants = [skill] + ALIASES.get(skill, [])
    for variant in variants:
        pattern = r"(?<!\w)" + re.escape(variant) + r"(?!\w)"
        if re.search(pattern, text or "", re.IGNORECASE):
            return True
    return False


def find_skills(text: str, catalogue: list[str]) -> list[str]:
    """Catalogue skills present in text, in catalogue order, no duplicates."""
    seen: list[str] = []
    for skill in catalogue:
        if skill not in seen and skill_hit(text, skill):
            seen.append(skill)
    return seen

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?:\+?91[\s\-]?)?[6-9]\d{9}")
_EXP_RE = re.compile(r"(\d+(?:\.\d+)?)\s*\+?\s*years?(?:\s+of)?\s+experience", re.IGNORECASE)
_EXP_ALT_RE = re.compile(r"experience\s*(?:of|:)?\s*(\d+(?:\.\d+)?)\s*\+?\s*years?", re.IGNORECASE)
_CTC_RE = re.compile(
    r"(current|expected|present)\s*ctc\s*(?:is|:)?\s*(?:rs\.?|inr|₹)?\s*([\d,]+(?:\.\d+)?)\s*(lpa|l|lakh|k)?",
    re.IGNORECASE,
)
_CTC_SHORT_RE = re.compile(r"\b([\d]+(?:\.\d+)?)\s*lpa\b", re.IGNORECASE)
_NOTICE_RE = re.compile(r"notice\s*period\s*(?:is|:)?\s*(\d+)\s*(day|month)", re.IGNORECASE)
_IMMEDIATE_RE = re.compile(r"immediate\s+join", re.IGNORECASE)
_DEGREE_RES = [
    re.compile(r"\bB\.?\s*Tech\b", re.IGNORECASE),
    re.compile(r"\bM\.?\s*Tech\b", re.IGNORECASE),
    re.compile(r"\bBCA\b", re.IGNORECASE),
    re.compile(r"\bMCA\b", re.IGNORECASE),
    re.compile(r"\bMBA\b", re.IGNORECASE),
    re.compile(r"\bB\.?\s*Sc\b", re.IGNORECASE),
    re.compile(r"\bM\.?\s*Sc\b", re.IGNORECASE),
    re.compile(r"\bPh\.?\s*D\b", re.IGNORECASE),
    re.compile(r"\bBachelor(?:'s)?\s+of\s+\w+", re.IGNORECASE),
    re.compile(r"\bMaster(?:'s)?\s+of\s+\w+", re.IGNORECASE),
]


def _lpa(value: float, unit: str | None) -> float:
    if not unit:
        return value
    u = unit.lower()
    if u in ("lpa", "l", "lakh"):
        return value
    if u == "k":
        return round(value / 100, 2)
    return value


def extract_text(data: bytes, kind: str) -> str:
    """Extract raw text from validated upload bytes (deterministic)."""
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    if kind == "doc":
        return _extract_doc_strings(data)
    raise ValueError(f"Unsupported kind '{kind}'.")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) > 20:
        raise ValueError("Resume exceeds the 20-page limit.")
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    if len(doc.paragraphs) > 2000:
        raise ValueError("Resume is too long.")
    return "\n".join(p.text for p in doc.paragraphs).strip()


def _extract_doc_strings(data: bytes) -> str:
    """Best-effort printable-strings extraction for legacy .doc (lossy)."""
    # Documented limitation: legacy .doc has no pure-Python reader here;
    # strings extraction recovers emails/phones/skills but not layout.
    text = re.sub(r"[^\x20-\x7e\n\r\t]", " ", data.decode("latin-1", errors="ignore"))
    chunks = [c.strip() for c in re.split(r"\s{2,}|\n+", text) if c.strip()]
    return "\n".join(chunks[:500])


def deterministic_parse(raw_text: str) -> dict:
    """Structure raw resume text into ParsedResume-compatible dict."""
    text = raw_text or ""
    email_m = _EMAIL_RE.search(text)
    phone_m = _PHONE_RE.search(text)
    email = email_m.group(0) if email_m else ""
    phone = re.sub(r"[\s\-]", "", phone_m.group(0)) if phone_m else ""

    lowered = text.lower()
    skills = find_skills(text, SKILLS)
    languages = find_skills(text, LANGUAGES)
    technologies = [s for s in skills if s not in languages]

    exp: float | None = None
    for rx in (_EXP_RE, _EXP_ALT_RE):
        m = rx.search(text)
        if m:
            try:
                exp = float(m.group(1))
            except ValueError:
                exp = None
            break

    education = sorted({m.group(0).strip() for rx in _DEGREE_RES for m in rx.finditer(text)})

    ctc_cur: float | None = None
    ctc_exp: float | None = None
    for m in _CTC_RE.finditer(text):
        try:
            val = _lpa(float(m.group(2).replace(",", "")), m.group(3))
        except ValueError:
            continue
        if m.group(1).lower() in ("current", "present") and ctc_cur is None:
            ctc_cur = val
        elif m.group(1).lower() == "expected" and ctc_exp is None:
            ctc_exp = val
    if ctc_cur is None and ctc_exp is None:
        shorts = [float(m.group(1)) for m in _CTC_SHORT_RE.finditer(text)]
        if shorts:
            ctc_cur = shorts[0]

    notice: int | None = None
    if _IMMEDIATE_RE.search(text):
        notice = 0
    else:
        m = _NOTICE_RE.search(text)
        if m:
            n = int(m.group(1))
            notice = n * 30 if m.group(2).lower().startswith("month") else n

    name = ""
    for line in (ln.strip() for ln in text.splitlines() if ln.strip()):
        if "@" in line or re.search(r"\d", line) or len(line) > 60:
            continue
        words = line.split()
        if 2 <= len(words) <= 5 and all(re.match(r"^[A-Z][\w.'-]*$", w) for w in words):
            name = line
            break

    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    summary = " ".join(lines[1:4])[:500] if len(lines) > 1 else text[:500]

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "skills": skills,
        "languages": languages,
        "technologies": technologies,
        "experienceYears": exp,
        "education": education,
        "ctcCurrentLpa": ctc_cur,
        "ctcExpectedLpa": ctc_exp,
        "noticePeriodDays": notice,
        "summary": summary,
    }
