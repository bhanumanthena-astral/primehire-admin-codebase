"""JD template + parse/map/validate for job creation (PRD Jobs §9-§11).

Prescribed template format (documented assumption — the PRD file itself is
absent from the workspace; the message spec fixes the labels and the error
strings, not the file layout): a DOCX whose body carries one ``Label: value``
line per field below, in order, with the Job Description running from its
label to end-of-document.

Flow: Upload → Extract (isolated child, reused) → Map → Review → Save.
This module is stateless: nothing is persisted here. The reviewed mapping
feeds the normal POST /api/jobs create endpoint (Slice 2), which owns
persistence, assignee checks, and notifications.
"""

from __future__ import annotations

import io
import logging
import re
from datetime import datetime, timezone
from typing import Any

logger = logging.getLogger(__name__)

JD_TEMPLATE_VERSION = "v1"

#: Exact labels, in template order. Matching is case-insensitive.
TEMPLATE_LABELS = (
    "Company Name",
    "Job Role",
    "Job Title",
    "Minimum Experience",
    "Maximum Experience",
    "Number of Positions",
    "Keywords",
    "Department",
    "Opened At",
    "Closes At",
    "Assignee",
    "Work Mode",
    "Location",
    "Job Description",
)

#: Labels whose VALUES must be non-empty (PRD §1 mandatory fields).
#: Keywords is recommended-only; Work Mode/Location are optional;
#: Opened At defaults to now (system source) when absent.
MANDATORY_LABELS = frozenset({
    "Company Name",
    "Job Role",
    "Job Title",
    "Minimum Experience",
    "Maximum Experience",
    "Number of Positions",
    "Department",
    "Closes At",
    "Assignee",
    "Job Description",
})

INVALID_TEMPLATE_MESSAGE = (
    "Invalid JD template. Please download the latest template "
    "and upload the completed template."
)
MISSING_FIELDS_MESSAGE = (
    "Please complete all mandatory fields before uploading the JD."
)

_LABEL_TO_KEY = {
    "company name": "companyName",
    "job role": "jobRole",
    "job title": "title",
    "minimum experience": "minExperienceYears",
    "maximum experience": "maxExperienceYears",
    "number of positions": "positionsTotal",
    "keywords": "keywords",
    "department": "department",
    "opened at": "openedAt",
    "closes at": "closesAt",
    "assignee": "assignee",
    "work mode": "workMode",
    "location": "location",
    "job description": "jdHtml",
}

_DATE_FORMATS = (
    "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S%z",
    "%d-%m-%Y", "%d/%m/%Y", "%d %b %Y", "%d %B %Y",
    "%b %d, %Y", "%B %d, %Y", "%m/%d/%Y",
)

_WORK_MODES = {
    "onsite": "ONSITE", "on-site": "ONSITE", "on site": "ONSITE", "office": "ONSITE",
    "remote": "REMOTE", "wfo": "ONSITE", "wfh": "REMOTE",
    "hybrid": "HYBRID",
}


def build_template_docx() -> bytes:
    """Build the prescribed blank JD template (v1) as DOCX bytes."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    heading = doc.add_heading("Job Description Template (v1)", level=1)
    for run in heading.runs:
        run.font.size = Pt(16)
    doc.add_paragraph(
        "Fill in every mandatory field, then upload this file while creating "
        "a job (Upload JD). Do not rename or remove the labels."
    )
    placeholders = {
        "Company Name": "Acme Pvt Ltd",
        "Job Role": "Backend Engineer",
        "Job Title": "Senior Backend Engineer",
        "Minimum Experience": "2",
        "Maximum Experience": "5",
        "Number of Positions": "3",
        "Keywords": "Python, SQL, AWS",
        "Department": "Engineering",
        "Opened At": "2026-10-01",
        "Closes At": "2026-12-31",
        "Assignee": "hr@company.com",
        "Work Mode": "Hybrid",
        "Location": "Hyderabad",
        "Job Description": "Describe the role, responsibilities and must-have skills here.",
    }
    for label in TEMPLATE_LABELS:
        doc.add_paragraph(f"{label}: {placeholders[label]}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _split_label_line(line: str) -> tuple[str, str] | None:
    if ":" not in line:
        return None
    label, _, value = line.partition(":")
    label = label.strip()
    if not label:
        return None
    return label, value.strip()


def parse_template_text(text: str) -> tuple[dict[str, str], list[str]]:
    """Map ``Label: value`` lines to raw values.

    Returns (values_by_label, missing_labels). The Job Description value is
    every line after its label through end-of-document.
    """
    found: dict[str, list[str]] = {}
    seen: set[str] = set()
    last: str | None = None
    in_jd = False
    jd_lines: list[str] = []
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if in_jd:
            jd_lines.append(line)
            continue
        split = _split_label_line(line)
        if split is not None:
            canonical: str | None = None
            for candidate in TEMPLATE_LABELS:
                if candidate.lower() == split[0].lower():
                    canonical = candidate
                    break
            if canonical is None:
                # Unknown label: treat as a wrapped continuation of the
                # previous value rather than dropping user content.
                if last is not None:
                    found[last].append(line)
                continue
            seen.add(canonical)
            if canonical == "Job Description":
                in_jd = True
                if split[1]:
                    jd_lines.append(split[1])
                last = None
                continue
            found.setdefault(canonical, []).append(split[1])
            last = canonical
        elif last is not None:
            found[last].append(line)
    values = {label: "\n".join(parts).strip() for label, parts in found.items()}
    if jd_lines:
        values["Job Description"] = "\n".join(jd_lines).strip()
    missing = [label for label in TEMPLATE_LABELS if label not in seen]
    return values, missing


def _parse_number(value: str, field: str, errors: list[dict[str, str]]) -> float | None:
    match = re.search(r"-?\d+(?:\.\d+)?", (value or "").strip())
    if not match:
        errors.append({"field": field, "message": f"{field} must be a number"})
        return None
    return float(match.group(0))


def _parse_positions(value: str, errors: list[dict[str, str]]) -> int | None:
    text = (value or "").strip()
    if not re.fullmatch(r"\d+", text):
        errors.append({"field": "positionsTotal",
                       "message": "positionsTotal must be a positive integer"})
        return None
    number = int(text)
    if number < 1 or number > 1000:
        errors.append({"field": "positionsTotal",
                       "message": "positionsTotal must be between 1 and 1000"})
        return None
    return number


def _parse_date(value: str, field: str, errors: list[dict[str, str]]) -> datetime | None:
    text = (value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        pass
    for fmt in _DATE_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    errors.append({"field": field, "message": f"{field} must be a valid date"})
    return None


async def map_template_values(
    values: dict[str, str], db: Any, org_id: str
) -> tuple[dict[str, Any], list[dict[str, str]]]:
    """Convert raw label values to a JobCreate-shaped mapping.

    Returns (mapped, field_errors). Assignee accepts an email or a userId and
    is resolved + active-checked here; failures are field errors, never fatal.
    Structural/Python-level checks (ranges, caps, dates order, JD length) are
    enforced by constructing JobCreate — a single source of validation truth.
    """
    from ..models.user import UserRepository
    from ..schemas.hiring import JobCreate, normalize_keywords

    errors: list[dict[str, str]] = []
    mapped: dict[str, Any] = {}

    for label, key in (("Company Name", "companyName"), ("Job Role", "jobRole"),
                       ("Job Title", "title"), ("Department", "department"),
                       ("Location", "location")):
        raw = (values.get(label) or "").strip()
        if raw:
            mapped[key] = raw

    min_exp = _parse_number(values.get("Minimum Experience", ""), "minExperienceYears", errors) \
        if (values.get("Minimum Experience") or "").strip() else None
    max_exp = _parse_number(values.get("Maximum Experience", ""), "maxExperienceYears", errors) \
        if (values.get("Maximum Experience") or "").strip() else None
    if min_exp is not None:
        mapped["minExperienceYears"] = min_exp
    if max_exp is not None:
        mapped["maxExperienceYears"] = max_exp

    positions_raw = (values.get("Number of Positions") or "").strip()
    if positions_raw:
        positions = _parse_positions(positions_raw, errors)
        if positions is not None:
            mapped["positionsTotal"] = positions

    keywords_raw = (values.get("Keywords") or "").strip()
    if keywords_raw:
        try:
            mapped["keywords"] = normalize_keywords(re.split(r"[,;\n]+", keywords_raw))
        except ValueError as exc:
            errors.append({"field": "keywords", "message": str(exc)})

    opened = _parse_date(values.get("Opened At", ""), "openedAt", errors) \
        if (values.get("Opened At") or "").strip() else None
    closes = _parse_date(values.get("Closes At", ""), "closesAt", errors) \
        if (values.get("Closes At") or "").strip() else None
    mapped["openedAt"] = opened or datetime.now(timezone.utc)
    if closes is not None:
        mapped["closesAt"] = closes

    work_raw = (values.get("Work Mode") or "").strip().lower()
    if work_raw:
        mode = _WORK_MODES.get(work_raw)
        if mode is None:
            errors.append({"field": "workMode",
                           "message": "workMode must be one of ONSITE, REMOTE, HYBRID"})
        else:
            mapped["workMode"] = mode

    assignee_raw = (values.get("Assignee") or "").strip()
    assignee_uid: str | None = None
    if assignee_raw:
        repo = UserRepository(db)
        user = await repo.get_by_id(assignee_raw, org_id)
        if user is None and "@" in assignee_raw:
            user = await repo.get_by_email(assignee_raw.strip().lower(), org_id)
        if user is None:
            errors.append({"field": "assigneeUserId",
                           "message": "Assignee is not a valid user of this organization"})
        elif user.get("isActive", True) is False:
            errors.append({"field": "assigneeUserId",
                           "message": "Assignee is deactivated; choose an active user"})
        else:
            assignee_uid = str(user.get("userId"))
            mapped["assigneeUserId"] = assignee_uid
            if user.get("email"):
                mapped["assigneeEmail"] = user.get("email")

    jd_raw = (values.get("Job Description") or "").strip()
    if jd_raw:
        paras = "".join(f"<p>{p.strip()}</p>" for p in jd_raw.splitlines() if p.strip())
        mapped["jdHtml"] = paras or f"<p>{jd_raw}</p>"

    # Mandatory presence AND value validation share one truth: JobCreate
    # requires every PRD-mandatory field, so construction below reports each
    # missing/invalid field by name. (No separate presence pass — it would
    # shadow deeper errors and double-report fields.)
    # Single validation truth: the same schema the create endpoint enforces.
    # jobKey is assigned at save time (never in the template), so validate
    # with a placeholder and strip it from the reviewed mapping.
    # assigneeEmail is endpoint-snapshot state, not a create field.
    assignee_email = mapped.pop("assigneeEmail", None)
    try:
        validated = JobCreate(**{**mapped, "jobKey": "TEMPLATE-CHECK"})
    except Exception as exc:  # noqa: BLE001 — translate pydantic errors to fields
        for err in getattr(exc, "errors", lambda: [])():
            loc = err.get("loc", ())
            field = str(loc[0]) if loc else _infer_field(str(err.get("msg", "")))
            if not any(e["field"] == field for e in errors):
                errors.append({"field": field, "message": str(err.get("msg", "invalid value"))})
        return {}, errors

    out = validated.model_dump(mode="json")
    out.pop("jobKey", None)
    if assignee_uid and assignee_email:
        out["assigneeEmail"] = assignee_email
    return out, []


def _infer_field(message: str) -> str:
    """Attribute a model-level validation message to the closest field."""
    lowered = message.lower()
    for field in ("maxExperienceYears", "minExperienceYears", "closesAt",
                  "openedAt", "positionsTotal", "keywords", "assigneeUserId"):
        if field.lower() in lowered:
            return field
    if "job description" in lowered:
        return "jdHtml"
    return "job"
